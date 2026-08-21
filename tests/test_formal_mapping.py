import csv
from pathlib import Path

from openpyxl import load_workbook

from scripts.formalize_qbc_mapping import (
    APPROVAL_HEADERS,
    approval_rows,
    merge_approval_rows,
)


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "outputs" / "qbc_ig_mapping" / "QBC_FHIR_Mapping_TaskSpec_v1.0-preview.1.xlsx"
FORMAL = ROOT / "outputs" / "qbc_ig_mapping" / "qbc_fhir_formal_mapping.csv"
APPROVALS = ROOT / "outputs" / "qbc_ig_mapping" / "qbc_mapping_approval_register.csv"
BUILD_SCRIPT = ROOT / "scripts" / "build_qbc_ig_mapping.py"


def rows(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def test_formal_mapping_has_one_complete_deterministic_rule_per_qbc_tag():
    mapping = rows(FORMAL)
    assert len(mapping) == 115
    assert len({row["QBC Tag"] for row in mapping}) == 115
    assert {row["Mapping Rule ID"] for row in mapping} == {f"QBC-MAP-{row['QBC Tag']}" for row in mapping}
    required = {
        "Target Resource", "Target Element", "Target datatype", "Target cardinality",
        "Target profile canonical", "Target profile version", "Transform rule", "Null handling",
        "Repeat handling", "Information-loss policy", "Round-trip policy", "Technical status",
        "Approval gates", "Release status", "Current bridge input", "Target source evidence",
        "Canonical fact", "Projection role",
    }
    for row in mapping:
        assert all(row[name].strip() for name in required), (row["QBC Tag"], row)
        assert " / " not in row["Target Resource"]
        assert " / " not in row["Target Element"]
        assert row["Technical status"] == "implemented-and-tested"


def test_profile_and_semantic_versions_are_pinned():
    for row in rows(FORMAL):
        profile = row["Target profile canonical"]
        expected = "1.0.0-preview.1" if "TW-Breast-Cancer-FHIR-IG/StructureDefinition" in profile else "1.0.0" if "twcore.mohw.gov.tw" in profile else "4.0.1"
        assert row["Target profile version"] == expected
        if row["Semantic alignment canonical"]:
            assert row["Semantic alignment version"] == "4.0.0"


def test_known_lossy_mappings_fail_closed_instead_of_inference():
    mapping = {row["QBC Tag"]: row for row in rows(FORMAL)}
    for tag in ["D023", "D057", "D085"]:
        assert "do not infer" in mapping[tag]["Information-loss policy"]
        assert "83052-1" in mapping[tag]["Terminology/ConceptMap"]
    for tag in ["D002", "D026", "D059"]:
        assert "lossy unless actual Organization" in mapping[tag]["Information-loss policy"]


def test_approval_register_contains_every_release_authority_and_no_fake_signature():
    approvals = rows(APPROVALS)
    ids = {row["Gate ID"] for row in approvals}
    assert {
        "QBC-TM02-SURGERY-POSTOP-REQUIRED", "QBC-DIAGTYPE3-RECURRENCE-REQUIRED",
        "QBC-XML-TABLE1-CLOSING-TAGS", "QBC-FAQ-PREOP-HORMONE-3D-DIRECT-SURGERY",
        "FHIR-MCODE-SCOPE", "FHIR-EXTENSIONS", "TERM-PR", "TERM-HER2-FISH", "TERM-PDL1",
        "GOV-AJCC", "GOV-SNOMED-DRUG", "SEC-PRIVACY", "UAT-VPN", "PUB-RELEASE",
    } == ids
    assert all(row["Status"] != "approved" for row in approvals)
    assert all(row["Required signer"].strip() and row["Acceptance evidence"].strip() for row in approvals)


def test_mapping_rebuild_preserves_a_decision_only_when_context_is_unchanged():
    first = dict(zip(APPROVAL_HEADERS[:6], approval_rows()[0]))
    first.update({field: "" for field in APPROVAL_HEADERS[6:]})
    first.update({
        "Status": "approved",
        "Decision (approve/reject/revise)": "approve",
        "Signer name": "Authorized owner",
        "Signer organization/title": "Hospital / QBC owner",
        "Decision date": "2026-08-21",
        "Evidence URI/path": "evidence/qbc-rule.json",
        "Signed artifact SHA-256": "a" * 64,
        "Notes": "signed decision",
    })
    merged = merge_approval_rows([first])
    preserved = dict(zip(APPROVAL_HEADERS, merged[0]))
    assert preserved["Status"] == "approved"
    assert preserved["Signer name"] == "Authorized owner"
    assert preserved["Signed artifact SHA-256"] == "a" * 64

    first["Proposed decision"] += " changed"
    invalidated = dict(zip(APPROVAL_HEADERS, merge_approval_rows([first])[0]))
    assert invalidated["Status"] == "pending-human-signoff"
    assert invalidated["Decision (approve/reject/revise)"] == ""
    assert invalidated["Signer name"] == ""
    assert invalidated["Signed artifact SHA-256"] == ""


def test_workbook_preserves_design_sheet_and_adds_formal_and_approval_sheets():
    wb = load_workbook(BOOK, read_only=True, data_only=False)
    assert "Field_Mapping_115" in wb.sheetnames
    assert "Formal_Mapping_115" in wb.sheetnames
    assert "Approval_Register" in wb.sheetnames
    assert "Architecture" in wb.sheetnames
    assert "Data_Flow" in wb.sheetnames
    assert sum(1 for ws in wb for row in ws.iter_rows() for cell in row if cell.data_type == "f") == 0


def test_fsh_defines_all_approved_candidate_artifact_types():
    fsh = (ROOT / "ig" / "input" / "fsh" / "extensions.fsh").read_text(encoding="utf-8")
    assert fsh.count("\nExtension:") == 9
    assert "Instance: QBCGenderToFHIRAdministrativeGender" in fsh
    assert "Instance: QBCLateralityToSNOMEDCT" in fsh
    assert "CodeSystem: QBCWorkflowCodeSystem" in fsh
    assert "LOINC" not in fsh  # assay-specific LOINC decisions stay governed in the mapping register


def test_breast_cancer_common_layer_owns_qbc_task():
    config = (ROOT / "ig" / "sushi-config.yaml").read_text(encoding="utf-8")
    common = (ROOT / "ig" / "input" / "fsh" / "breast-common-profiles.fsh").read_text(encoding="utf-8")
    qbc = (ROOT / "ig" / "input" / "fsh" / "profiles.fsh").read_text(encoding="utf-8")
    assert "id: io.github.erichuang777777.breast-cancer" in config
    assert "version: 1.0.0-preview.1" in config
    assert "Profile: BreastCancerPatient" in common
    assert "Profile: BreastCancerPrimaryCondition" in common
    assert "Profile: BreastCancerTumorMarkerObservation" in common
    assert "Profile: BreastCancerSourceDiagnosticReport" in common
    assert "Profile: BreastCancerSourceObservation" in common
    assert "Parent: BreastCancerSourceObservation" in common
    assert "Profile: QBCPatient\nParent: BreastCancerPatient" in qbc


def test_mapping_distinguishes_bridge_input_from_target_source_evidence():
    mapping = {row["QBC Tag"]: row for row in rows(FORMAL)}
    assert "secondary source" in mapping["D018"]["Current bridge input"]
    assert "非 Care Plan Task 輸出" in mapping["D018"]["Current bridge input"]
    assert "pathology/laboratory" in mapping["D018"]["Target source evidence"]
    assert "imaging" in mapping["D008"]["Target source evidence"]
    assert "Procedure" in mapping["TM02"]["Target source evidence"]
    assert all("derived projection" in row["Projection role"] for row in mapping.values())


def test_mapping_rebuild_uses_the_committed_versioned_extraction_not_a_private_docx():
    source = BUILD_SCRIPT.read_text(encoding="utf-8")
    assert '"qbc_workbench" / "data" / "qbc_fields.json"' in source
    assert "from docx import Document" not in source
    assert "Document(SPEC)" not in source
    mapping = rows(FORMAL)
    assert len({row["Source SHA-256"] for row in mapping}) == 1
    assert next(iter({row["Source SHA-256"] for row in mapping})) == (
        "f63be9eae07c552af4f03e9771e48bcda1c5b4eb14c36ae6291c1fc0fb4e069d"
    )
    for row in mapping:
        if len(row["QBC Tag"]) == 4 and row["QBC Tag"][0] == "D" and row["QBC Tag"][1:].isdigit():
            assert "secondary source" in row["Current bridge input"]
            assert "非 Care Plan Task 輸出" in row["Current bridge input"]
