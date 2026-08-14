from __future__ import annotations

import csv
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / "qbc_workbench" / "data" / "qbc_fields.json"
OUTDIR = ROOT / "outputs" / "qbc_ig_mapping"
FORMAL_CSV = OUTDIR / "qbc_fhir_formal_mapping.csv"
APPROVAL_CSV = OUTDIR / "qbc_mapping_approval_register.csv"
IG_PAGE = ROOT / "ig" / "input" / "pagecontent" / "formal-mapping.md"

FHIR = "http://hl7.org/fhir/StructureDefinition/"
TWCORE = "https://twcore.mohw.gov.tw/ig/twcore/StructureDefinition/"
QBC = "https://erichuang777777.github.io/TW-Breast-Cancer-FHIR-IG/StructureDefinition/"
MCODE = "http://hl7.org/fhir/us/mcode/StructureDefinition/"
PREVIEW_VERSION = "1.0.0-preview.1"

DATE_TAGS = {"BIRTHDAY", "P09", "D001", "D025", "D058", "TM09", "TM10", "T01", "T04", "T05", "T06"}
INTEGER_TAGS = {
    "P06", "TM01", "D015", "D017", "D021", "D039", "D040", "D042", "D043",
    "D045", "D046", "D049", "D051", "D055", "D071", "D072", "D074", "D075",
    "D077", "D079", "D083",
}
QUANTITY_TAGS = {"P03", "P04", "D015", "D017", "D021", "D047", "D049", "D051", "D055", "D077", "D079", "D083"}
FACILITY_TAGS = {"D002", "D026", "D059"}
HISTOLOGY_TAGS = {"D003", "D030"}
GRADE_TAGS = {"D004", "D031"}
T_TAGS = {"D005", "D032", "D060", "D064"}
N_TAGS = {"D006", "D033", "D061", "D065"}
M_TAGS = {"D007", "D034", "D062", "D066"}
STAGE_TAGS = {"D008", "D035", "D063", "D067"}
METASTASIS_TAGS = {"D009", "D036", "D068"}
METASTASIS_OTHER_TAGS = {"D010", "D037", "D069"}
QUAL_BIOMARKER_TAGS = {
    "D014", "D016", "D018", "D019", "D020", "D022", "D023", "D048", "D050",
    "D052", "D053", "D054", "D056", "D057", "D076", "D078", "D080", "D081",
    "D082", "D084", "D085",
}
PR_TAGS = {"D016", "D017", "D050", "D051", "D078", "D079"}
FISH_TAGS = {"D019", "D053", "D081"}
PDL1_TAGS = {"D023", "D057", "D085"}
NODE_STATUS_TAGS = {"D038", "D041", "D070", "D073"}
NODE_COUNT_TAGS = {"D039", "D040", "D042", "D043", "D045", "D046", "D071", "D072", "D074", "D075"}


def target(tag: str) -> dict[str, str]:
    data = {
        "resource": "Observation",
        "element": "Observation.valueString",
        "datatype": "string",
        "cardinality": "0..1",
        "profile": QBC + "qbc-raw-data-item-observation",
        "profile_version": PREVIEW_VERSION,
        "semantic_profile": "",
        "target_condition": "source value exists",
    }
    fixed = {
        "HOSPID": ("Organization", "Organization.identifier", "Identifier", "0..*", TWCORE + "Organization-twcore"),
        "ID": ("Patient", "Patient.identifier", "Identifier", "1..*", QBC + "qbc-patient"),
        "BIRTHDAY": ("Patient", "Patient.birthDate", "date", "1..1", QBC + "qbc-patient"),
        "DIAG_TYPE": ("EpisodeOfCare", "EpisodeOfCare.extension[qbc-enrollment-type]", "CodeableConcept", "0..1", FHIR + "EpisodeOfCare"),
        "LATERALITY": ("Condition", "Condition.bodySite", "CodeableConcept", "0..*", TWCORE + "Condition-twcore"),
        "P01": ("Patient", "Patient.name", "HumanName", "0..*", QBC + "qbc-patient"),
        "P02": ("Patient", "Patient.gender", "code", "0..1", QBC + "qbc-patient"),
        "P03": ("Observation", "Observation.valueQuantity", "Quantity", "1..1", TWCORE + "Observation-body-height-twcore"),
        "P04": ("Observation", "Observation.valueQuantity", "Quantity", "1..1", TWCORE + "Observation-body-weight-twcore"),
        "P05": ("Observation", "Observation.valueCodeableConcept", "CodeableConcept", "1..1", QBC + "qbc-data-item-observation"),
        "P06": ("Observation", "Observation.valueQuantity", "Quantity", "1..1", QBC + "qbc-quantity-data-item-observation"),
        "P07": ("Practitioner", "Practitioner.identifier", "Identifier", "0..*", TWCORE + "Practitioner-twcore"),
        "P08": ("EpisodeOfCare", "EpisodeOfCare.extension[qbc-case-class]", "CodeableConcept", "0..1", FHIR + "EpisodeOfCare"),
        "P09": ("EpisodeOfCare", "EpisodeOfCare.period.start", "dateTime", "0..1", FHIR + "EpisodeOfCare"),
        "D024": ("Observation", "Observation.valueCodeableConcept", "CodeableConcept", "1..1", QBC + "qbc-data-item-observation"),
        "D027": ("Procedure", "Procedure.code", "CodeableConcept", "1..1", TWCORE + "Procedure-twcore"),
        "D028": ("Observation", "Observation.valueCodeableConcept", "CodeableConcept", "1..1", QBC + "qbc-data-item-observation"),
        "D029": ("Procedure", "Procedure.code", "CodeableConcept", "1..1", TWCORE + "Procedure-twcore"),
        "D044": ("Observation", "Observation.valueBoolean", "boolean", "1..1", QBC + "qbc-data-item-observation"),
        "D047": ("Observation", "Observation.valueQuantity", "Quantity", "1..1", QBC + "qbc-tumor-size-observation"),
        "TM01": ("Procedure", "Procedure.extension[qbc-treatment-sequence]", "positiveInt", "0..1", TWCORE + "Procedure-twcore"),
        "TM02": ("Procedure", "Procedure.extension[qbc-treatment-type]", "CodeableConcept", "0..1", TWCORE + "Procedure-twcore"),
        "TM03": ("Procedure", "Procedure.extension[qbc-treatment-facility-class]", "CodeableConcept", "0..1", TWCORE + "Procedure-twcore"),
        "TM04": ("Procedure", "Procedure.code", "CodeableConcept", "1..1", TWCORE + "Procedure-twcore"),
        "TM05": ("MedicationAdministration", "MedicationAdministration.medicationCodeableConcept", "CodeableConcept", "1..1", FHIR + "MedicationAdministration"),
        "TM06": ("MedicationAdministration", "MedicationAdministration.medicationCodeableConcept.text", "string", "0..1", FHIR + "MedicationAdministration"),
        "TM07": ("Procedure", "Procedure.bodySite", "CodeableConcept", "0..*", TWCORE + "Procedure-twcore"),
        "TM08": ("Procedure", "Procedure.bodySite.text", "string", "0..1", TWCORE + "Procedure-twcore"),
        "TM09": ("Procedure", "Procedure.performedPeriod.start", "dateTime", "0..1", TWCORE + "Procedure-twcore"),
        "TM10": ("Procedure", "Procedure.performedPeriod.end", "dateTime", "0..1", TWCORE + "Procedure-twcore"),
        "T01": ("Observation", "Observation.effectiveDateTime", "dateTime", "0..1", QBC + "qbc-date-data-item-observation"),
        "T02": ("EpisodeOfCare", "EpisodeOfCare.extension[qbc-treatment-status]", "CodeableConcept", "0..1", FHIR + "EpisodeOfCare"),
        "T03": ("EpisodeOfCare", "EpisodeOfCare.extension[qbc-followup-status]", "CodeableConcept", "0..1", FHIR + "EpisodeOfCare"),
        "T04": ("EpisodeOfCare", "EpisodeOfCare.extension[qbc-transfer-date]", "date", "0..1", FHIR + "EpisodeOfCare"),
        "T05": ("EpisodeOfCare", "EpisodeOfCare.period.end", "dateTime", "0..1", FHIR + "EpisodeOfCare"),
        "T06": ("Patient", "Patient.deceasedDateTime", "dateTime", "0..1", QBC + "qbc-patient"),
    }
    if tag in fixed:
        resource, element, datatype, cardinality, profile = fixed[tag]
        data.update(resource=resource, element=element, datatype=datatype, cardinality=cardinality, profile=profile)
        return data
    if tag in {"D001", "D025", "D058"}:
        data.update(resource="Condition", element="Condition.onsetDateTime", datatype="dateTime", cardinality="0..1", profile=TWCORE + "Condition-twcore", semantic_profile=MCODE + "mcode-primary-cancer-condition")
    elif tag in FACILITY_TAGS:
        data.update(resource="Encounter", element="Encounter.extension[qbc-diagnosis-facility-class]", datatype="CodeableConcept", cardinality="0..1", profile=TWCORE + "Encounter-twcore")
    elif tag in HISTOLOGY_TAGS:
        data.update(resource="Condition", element="Condition.extension[histologyMorphologyBehavior]", datatype="CodeableConcept", cardinality="0..1", profile=TWCORE + "Condition-twcore", semantic_profile=MCODE + "mcode-primary-cancer-condition")
    elif tag in GRADE_TAGS:
        data.update(element="Observation.valueCodeableConcept", datatype="CodeableConcept", profile=QBC + "qbc-data-item-observation")
    elif tag in T_TAGS | N_TAGS | M_TAGS | STAGE_TAGS:
        slug = "tnm-primary-tumor-category" if tag in T_TAGS else "tnm-regional-nodes-category" if tag in N_TAGS else "tnm-distant-metastases-category" if tag in M_TAGS else "tnm-stage-group"
        data.update(element="Observation.valueCodeableConcept", datatype="CodeableConcept", profile=QBC + "qbc-data-item-observation", semantic_profile=MCODE + "mcode-" + slug)
    elif tag in METASTASIS_TAGS:
        data.update(resource="Condition", element="Condition.bodySite", datatype="CodeableConcept", cardinality="0..*", profile=TWCORE + "Condition-twcore", semantic_profile=MCODE + "mcode-secondary-cancer-condition")
    elif tag in METASTASIS_OTHER_TAGS:
        data.update(resource="Condition", element="Condition.bodySite.text", datatype="string", profile=TWCORE + "Condition-twcore", semantic_profile=MCODE + "mcode-secondary-cancer-condition")
    elif tag == "D011":
        data.update(element="Observation.valueCodeableConcept", datatype="CodeableConcept", profile=QBC + "qbc-data-item-observation")
    elif tag == "D012":
        data.update(resource="Procedure", element="Procedure.code", datatype="CodeableConcept", cardinality="1..1", profile=TWCORE + "Procedure-twcore")
    elif tag == "D013":
        data.update(element="Observation.valueCodeableConcept", datatype="CodeableConcept", profile=QBC + "qbc-data-item-observation")
    elif tag in QUAL_BIOMARKER_TAGS | NODE_STATUS_TAGS:
        data.update(element="Observation.valueCodeableConcept", datatype="CodeableConcept", profile=QBC + "qbc-data-item-observation", semantic_profile=MCODE + "mcode-tumor-marker-test" if tag not in {"D022", "D056", "D084"} else MCODE + "mcode-genomic-variant")
    elif tag in QUANTITY_TAGS:
        data.update(element="Observation.valueQuantity", datatype="Quantity", profile=QBC + "qbc-quantity-data-item-observation", semantic_profile=MCODE + "mcode-tumor-marker-test")
    elif tag in NODE_COUNT_TAGS:
        data.update(element="Observation.valueInteger", datatype="integer", profile=QBC + "qbc-integer-data-item-observation")
    return data


def terminology(tag: str) -> tuple[str, str, str]:
    if tag == "P02":
        return "QBCGenderToFHIRAdministrativeGender", "equivalent", "QBC 0/1/2/3 → male/female/other/unknown"
    if tag == "LATERALITY":
        return "QBCLateralityToSNOMEDCT", "equivalent", "L/R → SNOMED CT 7771000/24028007"
    if tag in {"P03", "P04"}:
        return "LOINC 8302-2|29463-7 + UCUM cm|kg", "equivalent", "decimal → Quantity"
    if tag in {"D014", "D015", "D048", "D049", "D076", "D077"}:
        return "LOINC 85329-1 + UCUM %", "related-to", "qualitative and percentage values remain separate QBC observations"
    if tag in PR_TAGS:
        return "LOINC 85339-0 (presence); 85325-9 (%)", "related-to", "assay=immune stain; preserve original QBC value"
    if tag in FISH_TAGS:
        return "LOINC 85318-4 (breast specimen FISH); 31150-6 fallback", "related-to", "use 85318-4 only when breast specimen and FISH are confirmed"
    if tag in PDL1_TAGS:
        return "No fixed LOINC without clone/assay; candidates 83052-1, 83055-4, 83057-0", "unmatched", "QBC lacks clone, CPS/TPS and threshold; preserve raw value and method metadata"
    if tag in T_TAGS | N_TAGS | M_TAGS | STAGE_TAGS:
        return "AJCC TNM/Stage (edition required)", "narrower", "QBC coarsens TNM; preserve original code and AJCC edition"
    if tag in DATE_TAGS:
        return "FHIR date/dateTime", "equivalent", "YYYYMMDD → YYYY-MM-DD; timezone not inferred"
    return "QBC field-specific ValueSet or source terminology", "equivalent", "copy or field-specific ConceptMap; raw QBC value retained in Provenance"


def transform(tag: str, spec: dict) -> str:
    if tag == "P02":
        return "translate(P02, QBCGenderToFHIRAdministrativeGender)"
    if tag == "LATERALITY":
        return "translate(LATERALITY, QBCLateralityToSNOMEDCT)"
    if tag in DATE_TAGS:
        return "parse YYYYMMDD as FHIR date/dateTime; reject invalid calendar dates"
    if tag in QUANTITY_TAGS:
        unit = "cm" if tag in {"P03", "D047"} else "kg" if tag == "P04" else "%"
        return f"parse decimal; set UCUM system and code '{unit}'"
    if spec.get("multi_select"):
        return "split on ASCII comma; trim; de-duplicate; preserve source order"
    if spec.get("allowed_values"):
        return "validate against QBC field ValueSet; translate when ConceptMap exists; preserve raw code"
    return "copy after Big5/max-length validation; preserve exact source in Provenance"


def approvals_for(tag: str, semantic_profile: str) -> list[str]:
    gates = []
    if tag in {"D027", "D028", "D029", "D030", "D031", "D032", "D033", "D034", "D035", "D036", "D037", "D038", "D039", "D040", "D041", "D042", "D043", "D044", "D045", "D046", "D047", "D048", "D049", "D050", "D051", "D052", "D053", "D054", "D055", "D056", "D057"}:
        gates.append("QBC-TM02-SURGERY-POSTOP-REQUIRED")
    if tag.startswith("D") and tag[1:].isdigit() and 58 <= int(tag[1:]) <= 85:
        gates.append("QBC-DIAGTYPE3-RECURRENCE-REQUIRED")
    if tag in PR_TAGS:
        gates.append("TERM-PR")
    if tag in FISH_TAGS:
        gates.append("TERM-HER2-FISH")
    if tag in PDL1_TAGS:
        gates.append("TERM-PDL1")
    if tag in T_TAGS | N_TAGS | M_TAGS | STAGE_TAGS:
        gates.append("GOV-AJCC")
    if semantic_profile:
        gates.append("FHIR-MCODE-SCOPE")
    if tag in {"DIAG_TYPE", "P08", "D002", "D026", "D059", "TM01", "TM02", "TM03", "T02", "T03", "T04"}:
        gates.append("FHIR-EXTENSIONS")
    return gates


def lineage(tag: str) -> tuple[str, str, str, str]:
    current = "癌症診療計畫書（目前 bridge input）"
    target_source = "signed clinical source + provenance"
    canonical_fact = "BreastCancer canonical fact"
    if tag in {"HOSPID", "ID", "BIRTHDAY", "P01", "P02", "P03", "P04", "P05", "P06", "P07", "P08", "P09", "DIAG_TYPE", "LATERALITY"}:
        current = "癌症診療計畫書／病人、醫師或個管主檔"
        target_source = "patient/master data, clinical assessment or encounter source"
        canonical_fact = "Patient / EpisodeOfCare / Condition / Observation"
    elif tag.startswith("D"):
        if tag in HISTOLOGY_TAGS | GRADE_TAGS | QUAL_BIOMARKER_TAGS | QUANTITY_TAGS | NODE_STATUS_TAGS | NODE_COUNT_TAGS | {"D011", "D013", "D024", "D028", "D044"}:
            target_source = "pathology/laboratory DiagnosticReport + Observation + Specimen"
            canonical_fact = "pathology finding / tumor marker / specimen-linked observation"
        elif tag in T_TAGS | N_TAGS | M_TAGS | STAGE_TAGS | METASTASIS_TAGS | METASTASIS_OTHER_TAGS:
            target_source = "imaging, pathology and clinical staging evidence"
            canonical_fact = "TNM/stage or metastatic disease fact"
        else:
            target_source = "pathology, imaging or signed clinical DiagnosticReport"
            canonical_fact = "diagnosis or treatment-related clinical fact"
    elif tag.startswith("TM"):
        current = "癌症診療計畫書／個管整理的治療資料"
        target_source = "Procedure, MedicationAdministration, radiotherapy or treatment-system source"
        canonical_fact = "cancer treatment fact"
    elif tag.startswith("T0"):
        current = "癌症診療計畫書／個管追蹤資料"
        target_source = "follow-up encounter, vital status, transfer or outcome source"
        canonical_fact = "follow-up/outcome fact"
    return current, target_source, canonical_fact, "derived projection to QBC; plan and QBC are not source-of-truth facts"


def approval_rows() -> list[list[str]]:
    return [
        ["QBC-TM02-SURGERY-POSTOP-REQUIRED", "QBC business rule", "D027-D057 mandatory when TM02=1 and DIAG_TYPE in (1,2), field-specific conditions override", "QBC申報負責人", "pending-human-signoff", "official clarification or signed institutional decision + positive/negative UAT"],
        ["QBC-DIAGTYPE3-RECURRENCE-REQUIRED", "QBC business rule", "D058-D085 mandatory for DIAG_TYPE=3, field-specific conditions override", "QBC申報負責人", "pending-human-signoff", "signed decision + recurrence golden/negative cases"],
        ["QBC-XML-TABLE1-CLOSING-TAGS", "source interpretation", "Use well-formed XML and table-2 field definitions instead of malformed table-1 sample tags", "QBC申報負責人", "pending-human-signoff", "VPN test receipt for generated XML"],
        ["QBC-FAQ-PREOP-HORMONE-3D-DIRECT-SURGERY", "clinical classification", "Three days of preoperative hormone therapy remains DIAG_TYPE=1", "乳癌臨床專家", "pending-human-signoff", "signed clinical decision and sample case"],
        ["FHIR-MCODE-SCOPE", "FHIR architecture", "mCODE is semantic alignment only; no dual-conformance claim with TW Core", "FHIR實作負責人", "pending-human-signoff", "signed architecture decision; validate all cited canonical URLs against mCODE 4.0.0"],
        ["FHIR-EXTENSIONS", "FHIR architecture", "Approve contexts, datatypes and canonical URLs of QBC extensions", "FHIR實作負責人", "pending-human-signoff", "Publisher QA=0/0 and reviewed examples"],
        ["TERM-PR", "terminology", "Use LOINC 85339-0 for PR presence and 85325-9 for percentage when immune stain is confirmed", "病理／術語專家", "pending-human-signoff", "local LIS assay mapping and terminology license review"],
        ["TERM-HER2-FISH", "terminology", "Use LOINC 85318-4 for breast specimen FISH; otherwise 31150-6 only when tissue FISH is confirmed", "病理／術語專家", "pending-human-signoff", "local LIS specimen/method verification"],
        ["TERM-PDL1", "terminology", "Do not assign an assay-specific LOINC without clone/assay/CPS/TPS; preserve raw QBC result", "病理／術語專家", "pending-human-signoff", "local LIS captures clone, scoring method and threshold"],
        ["GOV-AJCC", "license/version governance", "Store AJCC edition and approve permitted use; QBC coarse stage remains source of truth for export", "癌登／法遵負責人", "pending-human-signoff", "edition decision and license evidence"],
        ["GOV-SNOMED-DRUG", "license/terminology governance", "Approve SNOMED CT and local drug-code ConceptMaps and licensing", "術語／法遵負責人", "pending-human-signoff", "licensed terminology release and reviewed ConceptMaps"],
        ["SEC-PRIVACY", "security/privacy", "Approve identifier handling, access control, audit, retention and de-identification", "個資／資安負責人", "pending-human-signoff", "DPIA/security review and penetration-test findings closed"],
        ["UAT-VPN", "acceptance", "Cross-system FHIR round-trip and official VPN XML acceptance", "QBC申報負責人＋系統負責人", "pending-external-acceptance", "receipts, error-code matrix and reconciliation report"],
        ["PUB-RELEASE", "publication", "Approve publisher identity, immutable canonical/package/version and promotion to active/non-experimental", "資料治理委員會／publisher", "blocked-until-other-gates-close", "all gates signed; release manifest and hashes"],
    ]


def style(ws, widths: dict[int, int]) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    for i, width in widths.items():
        ws.column_dimensions[get_column_letter(i)].width = width
    for c in ws[1]:
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(vertical="top", wrap_text=True)


def formalize(path: Path) -> Path:
    payload = json.loads(SPEC.read_text(encoding="utf-8"))
    specs = {f["tag"]: f for f in payload["fields"]}
    source = payload["source"]
    wb = load_workbook(path)
    source_ws = wb["Field_Mapping_115"]
    headers = [c.value for c in source_ws[1]]
    source_rows = {str(row[headers.index("QBC Tag")]): dict(zip(headers, row)) for row in source_ws.iter_rows(min_row=2, values_only=True)}

    for name in ["Architecture", "Data_Flow", "Formal_Mapping_115", "Approval_Register"]:
        if name in wb.sheetnames:
            del wb[name]
    formal_headers = [
        "Mapping Rule ID", "QBC Tag", "QBC 中文名稱", "Source specification", "Source SHA-256",
        "Source condition", "Source cardinality", "Source datatype", "Current bridge input", "Target source evidence",
        "Canonical fact", "Projection role", "Target Resource", "Target Element",
        "Target datatype", "Target cardinality", "Target profile canonical", "Target profile version",
        "Semantic alignment canonical", "Semantic alignment version", "Terminology/ConceptMap", "Map relationship", "Transform rule",
        "Null handling", "Repeat handling", "Information-loss policy", "Round-trip policy",
        "Technical status", "Approval gates", "Release status",
    ]
    ws = wb.create_sheet("Formal_Mapping_115", 2)
    ws.append(formal_headers)
    formal_rows = []
    for tag, spec in specs.items():
        trg = target(tag)
        current_input, target_evidence, canonical_fact, projection_role = lineage(tag)
        profile_version = PREVIEW_VERSION if trg["profile"].startswith(QBC) else "1.0.0" if trg["profile"].startswith(TWCORE) else "4.0.1"
        term, relation, term_note = terminology(tag)
        gates = approvals_for(tag, trg["semantic_profile"])
        required = "1..*" if spec.get("always_required") and tag.startswith(("TM", "T0")) else "0..*" if tag.startswith(("TM", "T0")) else "1..1" if spec.get("always_required") else "0..1"
        conditions = " OR ".join(spec.get("required_when") or []) or ("true" if spec.get("always_required") else "source value exists")
        source_type = spec.get("data_type") or "string"
        if spec.get("multi_select"):
            source_type += " (comma-separated repeating codes)"
        loss = "lossless with raw QBC value retained in Provenance"
        if tag in FACILITY_TAGS:
            loss = "lossy unless actual Organization is also supplied; 1/2 only records same/other facility"
        elif tag in PDL1_TAGS:
            loss = "lossy for assay/clone/CPS/TPS because QBC does not carry them; do not infer"
        elif tag in T_TAGS | N_TAGS | M_TAGS | STAGE_TAGS:
            loss = "potentially lossy toward mCODE/AJCC granularity; retain exact QBC stage and edition"
        row = [
            f"QBC-MAP-{tag}", tag, source_rows[tag]["QBC 中文名稱"], source.get("version", ""), source.get("sha256", ""),
            conditions, required, source_type, current_input, target_evidence, canonical_fact, projection_role,
            trg["resource"], trg["element"], trg["datatype"], trg["cardinality"],
            trg["profile"], profile_version, trg["semantic_profile"], "4.0.0" if trg["semantic_profile"] else "", f"{term}; {term_note}", relation,
            transform(tag, spec), "reject missing when condition is true; otherwise omit target element; never invent a value",
            "create one target item per source repetition; multi-select values are separate codings unless profile requires one CodeableConcept",
            loss, "QBC export MUST use retained raw value plus validated canonical value; fail closed if they disagree",
            "implemented-and-tested", "; ".join(gates) or "none", "ready-for-human-review" if gates else "technically-ready",
        ]
        formal_rows.append(row)
        ws.append(row)
    style(ws, {1: 20, 2: 12, 3: 34, 4: 16, 5: 66, 6: 45, 7: 15, 8: 25,
               9: 42, 10: 55, 11: 45, 12: 50, 13: 25, 14: 52, 15: 19, 16: 17,
               17: 70, 18: 20, 19: 70, 20: 20, 21: 65, 22: 18, 23: 62, 24: 52,
               25: 52, 26: 58, 27: 60, 28: 22, 29: 48, 30: 24})
    ws.add_table(Table(displayName="QBCFormalMapping", ref=ws.dimensions))
    ws.tables["QBCFormalMapping"].tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)

    approval_headers = [
        "Gate ID", "Category", "Proposed decision", "Required signer", "Status", "Acceptance evidence",
        "Decision (approve/reject/revise)", "Signer name", "Signer organization/title", "Decision date",
        "Evidence URI/path", "Signed artifact SHA-256", "Notes",
    ]
    aws = wb.create_sheet("Approval_Register", 3)
    aws.append(approval_headers)
    for row in approval_rows():
        aws.append(row + [""] * 7)
    style(aws, {1: 34, 2: 25, 3: 90, 4: 34, 5: 30, 6: 80, 7: 28, 8: 24, 9: 34, 10: 18, 11: 55, 12: 66, 13: 50})
    aws.add_table(Table(displayName="QBCApprovalRegister", ref=aws.dimensions))
    aws.tables["QBCApprovalRegister"].tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)

    arch = wb.create_sheet("Architecture", 2)
    arch.append(["Layer", "Artifact owner", "Purpose", "Dependency/alignment", "Release status"])
    for row in [
        ["FHIR base", "HL7 International", "Base resources and exchange framework", "hl7.fhir.r4.core#4.0.1", "external published standard"],
        ["Taiwan base", "MOHW / TW Core", "Taiwan identifiers and reusable base profiles", "tw.gov.mohw.twcore#1.0.0", "external trial-use dependency"],
        ["Source evidence layer", "source systems / signed reports", "Pathology, laboratory, ultrasound, specimen and imaging lineage", "DiagnosticReport + Observation + Specimen + ImagingStudy", "target-state draft"],
        ["Breast cancer canonical facts", "Independent community draft", "Reusable diagnosis, stage, biomarker, treatment and outcome facts", "mCODE 4.0.0 and ICHOM Breast Cancer 1.0.0 semantic reference", "draft / experimental"],
        ["Derived document/collaboration layer", "clinical workflow", "Cancer diagnosis/treatment plan, treatment plan and MDT record", "generated from canonical facts; may be bridge input during migration", "partially modeled"],
        ["QBC/P4P task projection", "Independent community draft", "115-field mapping, QBC rules, reversible transforms and XML/VPN validation", "reads canonical facts or current plan bridge input", "1.0.0-preview.1"],
    ]:
        arch.append(row)
    style(arch, {1: 30, 2: 34, 3: 70, 4: 75, 5: 30})
    arch.add_table(Table(displayName="QBCArchitecture", ref=arch.dimensions))
    arch.tables["QBCArchitecture"].tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)

    flow = wb.create_sheet("Data_Flow", 3)
    flow.append(["State", "Input", "Transformation", "Canonical layer", "Output/task", "Governance rule"])
    for row in [
        ["Current bridge", "癌症診療計畫書", "document extraction + QBC review", "partial; gradually normalize to BreastCancer facts", "QBC FHIR Bundle → QBC XML", "plan is derived input, not original evidence"],
        ["Migration", "癌症診療計畫書", "bridge adapter + provenance + reconciliation", "BreastCancer canonical facts", "QBC and regenerated treatment-plan document", "retain document version/page and missing-original-evidence flag"],
        ["Target A", "pathology/laboratory/ultrasound/treatment sources", "source-specific mapping", "BreastCancer canonical facts", "generate cancer diagnosis/treatment plan", "signed source and specimen/imaging lineage retained"],
        ["Target B", "pathology/laboratory/ultrasound/treatment sources", "source-specific mapping", "BreastCancer canonical facts", "direct QBC projection", "same rules/version as plan-first pathway"],
        ["Target C", "BreastCancer canonical facts", "task-specific projection", "single reusable fact store", "drug prior authorization / registry / treatment plan / MDT", "task output never overwrites source facts silently"],
    ]:
        flow.append(row)
    style(flow, {1: 20, 2: 55, 3: 55, 4: 45, 5: 55, 6: 70})
    flow.add_table(Table(displayName="QBCDataFlow", ref=flow.dimensions))
    flow.tables["QBCDataFlow"].tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)

    readme = wb["README"]
    readme["B4"] = "1.0.0-preview.1（乳癌社群草稿中的 QBC／P4P Task；非官方、draft／experimental）"
    readme["A6"] = "正式化狀態"
    readme["B6"] = "Formal_Mapping_115 提供逐欄可執行規則；Approval_Register 保留正式／官方使用的治理與驗收閘門。社群 Preview 不等於權責核准。"
    readme["A7"] = "資料流狀態"
    readme["B7"] = "目前：癌症診療計畫書 → QBC Task。目標：原始報告／系統 → BreastCancer canonical facts → 計畫書或直接 QBC／其他 Task。"
    readme["A6"].font = copy(readme["A5"].font)
    readme["A6"].alignment = copy(readme["A5"].alignment)
    readme["A7"].font = copy(readme["A5"].font)
    readme["A7"].alignment = copy(readme["A5"].alignment)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(path)

    OUTDIR.mkdir(parents=True, exist_ok=True)
    with FORMAL_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(formal_headers)
        writer.writerows(formal_rows)
    with APPROVAL_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(approval_headers)
        writer.writerows(row + [""] * 7 for row in approval_rows())
    return path


if __name__ == "__main__":
    import sys

    print(formalize(Path(sys.argv[1]) if len(sys.argv) > 1 else OUTDIR / "QBC_FHIR_Mapping_TaskSpec_v1.0-preview.1.xlsx"))
