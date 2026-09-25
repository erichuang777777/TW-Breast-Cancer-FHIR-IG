from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXCLUDED_PARTS = {".git", ".pytest_cache", "__pycache__"}
ERROR_TOKENS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?"}


def files_under(root: Path):
    for path in root.rglob("*"):
        if path.is_file() and not EXCLUDED_PARTS.intersection(path.parts):
            yield path


def file_summary(groups: dict[str, Path]):
    rows = []
    for group, root in groups.items():
        counts = Counter()
        sizes = Counter()
        for path in files_under(root):
            ext = path.suffix.lower() or "[no extension]"
            counts[ext] += 1
            sizes[ext] += path.stat().st_size
        for ext in sorted(counts):
            rows.append([group, ext, counts[ext], sizes[ext]])
    return rows


def workbook_summary(groups: dict[str, Path]):
    profiles = Counter()
    variants = Counter()
    failures = Counter()
    for group, root in groups.items():
        for path in files_under(root):
            if path.suffix.lower() != ".xlsx" or path.name.startswith("~$"):
                continue
            try:
                wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
                workbook_shape = []
                for ws in wb.worksheets:
                    formula_count = 0
                    for row in ws.iter_rows():
                        formula_count += sum(
                            1 for cell in row if cell.data_type == "f"
                        )
                    profile = (
                        group,
                        ws.max_row,
                        ws.max_column,
                        formula_count,
                        ws.sheet_state,
                    )
                    profiles[profile] += 1
                    workbook_shape.append(profile[1:])
                signature = hashlib.sha256(
                    json.dumps(sorted(workbook_shape), ensure_ascii=False).encode("utf-8")
                ).hexdigest()[:12]
                variants[(group, signature, len(wb.sheetnames))] += 1
                wb.close()
            except Exception as exc:
                failures[(group, type(exc).__name__)] += 1

    profile_rows = [list(key) + [count] for key, count in sorted(profiles.items())]
    variant_rows = [list(key) + [count] for key, count in sorted(variants.items())]
    failure_rows = [list(key) + [count] for key, count in sorted(failures.items())]
    return profile_rows, variant_rows, failure_rows


def read_csv(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def field_manifest_rows(path: Path):
    manifest = json.loads(path.read_text(encoding="utf-8"))
    rows = []
    for mapping_id, entries in sorted(manifest.items()):
        for entry in entries:
            rows.append(
                [
                    mapping_id,
                    entry.get("report", ""),
                    entry.get("worksheet_column", ""),
                    entry.get("concept_zh", ""),
                    entry.get("source", ""),
                    entry.get("note", ""),
                ]
            )
    return rows


def add_sheet(wb: Workbook, title: str, headers: list[str], rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for cell in ws.iter_cols(min_col=index, max_col=index, max_row=ws.max_row):
            for item in cell:
                max_len = max(max_len, min(len(str(item.value or "")), 60))
        ws.column_dimensions[get_column_letter(index)].width = min(max_len + 2, 62)
    return ws


def build(args):
    groups = {
        "quality-source": Path(args.quality_root),
        "quarterly-source": Path(args.quarterly_root),
        "reference-implementation": Path(args.implementation_root),
    }
    missing = [label for label, root in groups.items() if not root.is_dir()]
    if missing:
        raise SystemExit(f"Missing source group(s): {', '.join(missing)}")

    repo = Path(args.repo_root).resolve()
    mapping_root = repo / "mappings" / "case-management"
    file_rows = file_summary(groups)
    profile_rows, variant_rows, failure_rows = workbook_summary(groups)
    manifest_rows = field_manifest_rows(groups["reference-implementation"] / "field_manifest.json")
    common_rows = read_csv(mapping_root / "breast-common-to-case-management.csv")
    criteria_rows = read_csv(mapping_root / "case-management-population-criteria.csv")
    measure_rows = read_csv(mapping_root / "case-management-measure-catalog.csv")

    common_headers = [
        "mapping_id", "concept", "target_profile", "target_path", "standard_code",
        "code_status", "report_family", "indicator_uses", "current_source",
        "current_source_kind", "fhir_gain", "fallback", "source_ids", "review_status",
    ]
    common_output = [[row.get(header, "") for header in common_headers] for row in common_rows]

    blockers = []
    for row in criteria_rows:
        if row.get("review_status") in {"blocking-data-gap", "open-question"} or row.get(
            "python_status"
        ) in {"not-implemented", "not-evaluable", "divergent"}:
            blockers.append(
                [
                    "criterion",
                    row.get("indicator_family", ""),
                    row.get("measure_id", ""),
                    row.get("criterion_id", ""),
                    row.get("criterion_zh", ""),
                    row.get("python_status", ""),
                    row.get("review_status", ""),
                    row.get("current_proxy", ""),
                    row.get("python_divergence", ""),
                ]
            )
    for row in common_rows:
        if row.get("review_status") == "blocking-data-gap":
            blockers.append(
                [
                    "mapping",
                    row.get("report_family", ""),
                    row.get("indicator_uses", ""),
                    row.get("mapping_id", ""),
                    row.get("concept", ""),
                    "",
                    row.get("review_status", ""),
                    row.get("current_source", ""),
                    row.get("fhir_gain", ""),
                ]
            )

    wb = Workbook()
    wb.remove(wb.active)
    notes = [
        ["產生時間", datetime.now().astimezone().isoformat(timespec="seconds")],
        ["用途", "品管＋季報 Task 的無個資來源結構盤點；不含檔名、病歷號、姓名或個案值。"],
        ["來源群組", "quality-source、quarterly-source、reference-implementation"],
        ["工作簿變體", "依 sheet 數、列欄尺寸、公式數與可見狀態建立結構雜湊；不使用 sheet 名稱或內容。"],
        ["欄位清冊", "來自 reference implementation 的 field_manifest.json，只列 schema label 與來源類型。"],
        ["FHIR 對應", "來自 repo mappings/case-management；此檔是盤點輸出，不自行核定臨床語意。"],
        ["隱私", "原始工作簿只在本機唯讀開啟；輸出不保存路徑、檔名、儲存格值或個案級統計。"],
    ]
    add_sheet(wb, "說明", ["項目", "內容"], notes)
    add_sheet(wb, "來源摘要", ["來源群組", "副檔名", "檔案數", "總位元組"], file_rows)
    add_sheet(
        wb,
        "工作簿變體",
        ["來源群組", "結構簽章", "sheet數", "工作簿數"],
        variant_rows,
    )
    add_sheet(
        wb,
        "工作表結構",
        ["來源群組", "列數", "欄數", "公式數", "顯示狀態", "出現次數"],
        profile_rows,
    )
    add_sheet(wb, "讀取失敗", ["來源群組", "錯誤類型", "檔案數"], failure_rows)
    add_sheet(
        wb,
        "欄位清冊",
        ["FHIR mapping id", "報表族", "工作表欄位", "共同概念", "來源類型", "備註"],
        manifest_rows,
    )
    add_sheet(wb, "FHIR對應覆蓋", common_headers, common_output)
    add_sheet(
        wb,
        "阻擋項",
        ["類型", "報表族", "Measure", "ID", "內容", "Python狀態", "審閱狀態", "目前代理", "差異／FHIR效益"],
        blockers,
    )
    measure_headers = list(measure_rows[0]) if measure_rows else []
    add_sheet(
        wb,
        "Measure目錄",
        measure_headers,
        [[row.get(header, "") for header in measure_headers] for row in measure_rows],
    )

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    verify = load_workbook(output, read_only=True, data_only=False)
    formulas = 0
    errors = []
    for ws in verify.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formulas += 1
                if isinstance(cell.value, str) and cell.value in ERROR_TOKENS:
                    errors.append(f"{ws.title}!{cell.coordinate}:{cell.value}")
    verify.close()
    if formulas or errors:
        raise SystemExit(f"Output verification failed: formulas={formulas}, errors={errors}")
    print(
        json.dumps(
            {
                "output": str(output),
                "source_groups": len(groups),
                "workbook_variants": len(variant_rows),
                "sheet_profiles": len(profile_rows),
                "field_manifest_rows": len(manifest_rows),
                "mapping_rows": len(common_output),
                "blocker_rows": len(blockers),
                "formula_count": formulas,
                "formula_errors": len(errors),
            },
            ensure_ascii=False,
        )
    )


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--quality-root", required=True)
    parser.add_argument("--quarterly-root", required=True)
    parser.add_argument("--implementation-root", required=True)
    parser.add_argument("--repo-root", default=str(Path(__file__).resolve().parents[1]))
    parser.add_argument(
        "--output",
        default=str(
            Path(__file__).resolve().parents[1]
            / "outputs"
            / "case_management_source_audit"
            / "case_management_source_inventory.xlsx"
        ),
    )
    return parser.parse_args()


if __name__ == "__main__":
    build(parse_args())
