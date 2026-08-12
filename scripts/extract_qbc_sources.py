from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pdfplumber
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "source_audit" / "extracted"
PDFS = [
    "1 全民健康保險乳癌照護品質提升方案_方案說明.pdf",
    "2 乳癌照護品質提升方案問答集(1150629修訂).pdf",
    "3 乳癌VPN系統說明會.pdf",
    "4 乳癌照護品質提升方案VPN使用者手冊.pdf",
    "5 批次上傳使用手冊_QBC_UploadXML_UserGuide.pdf",
]
DOCX = "6 批次上傳格式說明_QBC_乳癌照護品質提升方案_XML上傳_11507(定版).docx"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract_pdf(path: Path, output: Path) -> dict:
    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            chunks.extend([f"===== PAGE {page_number} =====", page.extract_text(layout=True) or "", ""])
        pages = len(pdf.pages)
    text = "\n".join(chunks)
    output.write_text(text, encoding="utf-8")
    return {"file": path.name, "type": "pdf", "sha256": sha256(path), "pages": pages, "extracted_lines": len(text.splitlines()), "output": str(output.relative_to(ROOT)).replace("\\", "/")}


def extract_docx(path: Path, output: Path) -> dict:
    document = Document(path)
    chunks: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        if paragraph.text.strip():
            chunks.append(f"PARAGRAPH {index}: {paragraph.text}")
    for table_number, table in enumerate(document.tables, 1):
        chunks.append(f"===== TABLE {table_number} =====")
        for row_number, row in enumerate(table.rows, 1):
            values = [" ".join(cell.text.split()) for cell in row.cells]
            chunks.append(f"ROW {row_number}: " + " | ".join(values))
    text = "\n".join(chunks) + "\n"
    output.write_text(text, encoding="utf-8")
    return {
        "file": path.name, "type": "docx", "sha256": sha256(path),
        "paragraphs": len(document.paragraphs), "tables": len(document.tables),
        "extracted_lines": len(text.splitlines()), "output": str(output.relative_to(ROOT)).replace("\\", "/"),
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = []
    for index, filename in enumerate(PDFS, 1):
        manifest.append(extract_pdf(ROOT / filename, OUT / f"{index:02d}_{Path(filename).stem}.txt"))
    manifest.append(extract_docx(ROOT / DOCX, OUT / f"06_{Path(DOCX).stem}.txt"))
    manifest_path = OUT.parent / "source_inventory.json"
    manifest_path.write_text(json.dumps({"sources": manifest}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(manifest_path)
    for item in manifest:
        print(f"{item['file']}: {item['extracted_lines']} lines")


if __name__ == "__main__":
    main()
