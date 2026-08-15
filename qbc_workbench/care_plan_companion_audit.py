from __future__ import annotations

from collections import Counter
from pathlib import Path
from statistics import median
from typing import Any
import json
import re
import unicodedata

from openpyxl import load_workbook
from pypdf import PdfReader


TOKEN_PATTERN = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.:/+-]{1,}|[\u3400-\u9fff]{2,}")


def _distribution(values: list[int | float]) -> dict[str, int | float]:
    if not values:
        return {"min": 0, "median": 0, "max": 0}
    return {
        "min": round(min(values), 4),
        "median": round(median(values), 4),
        "max": round(max(values), 4),
    }


def _tokens(values: list[str]) -> set[str]:
    text = unicodedata.normalize("NFKC", "\n".join(values)).lower()
    return {token for token in TOKEN_PATTERN.findall(text) if len(token) >= 2}


def _json_strings(value: Any) -> list[str]:
    strings: list[str] = []
    if isinstance(value, dict):
        for key, nested in value.items():
            strings.append(str(key))
            strings.extend(_json_strings(nested))
    elif isinstance(value, list):
        for nested in value:
            strings.extend(_json_strings(nested))
    elif isinstance(value, (str, int, float, bool)):
        strings.append(str(value))
    return strings


def _xlsx_strings(path: Path) -> tuple[list[str], dict[str, int]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values: list[str] = []
    nonempty_cells = 0
    formula_cells = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells += 1
                if cell.value is not None and str(cell.value).strip():
                    nonempty_cells += 1
                    values.append(str(cell.value))
    metrics = {
        "sheets": len(workbook.worksheets),
        "nonempty_cells": nonempty_cells,
        "formula_cells": formula_cells,
    }
    workbook.close()
    return values, metrics


def _pdf_strings(path: Path) -> tuple[list[str], dict[str, int]]:
    reader = PdfReader(path)
    texts: list[str] = []
    pages_with_text = 0
    extracted_characters = 0
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages_with_text += 1
            extracted_characters += len(text)
            texts.append(text)
    return texts, {
        "pages": len(reader.pages),
        "pages_with_text": pages_with_text,
        "extracted_characters": extracted_characters,
    }


def _coverage(source: set[str], reference: set[str]) -> float:
    return len(source & reference) / len(source) if source else 1.0


def _normalized_scalar(value: str) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).lower()


def _unmatched_cell_class(value: str, case_stem: str) -> str:
    normalized = _normalized_scalar(value)
    if _normalized_scalar(case_stem) in normalized:
        return "case-identifier-like"
    if re.search(r"(?:19|20)?\d{2}[/.-]\d{1,2}[/.-]\d{1,2}", normalized):
        return "date-like"
    if not re.search(r"[\u3400-\u9fff]", normalized) and re.fullmatch(
        r"[a-z0-9_.:/+\-(),%]+", normalized
    ):
        return "code-or-numeric-like"
    if re.search(r"[\u3400-\u9fff]", normalized) or len(normalized) >= 20:
        return "narrative-or-heading-like"
    return "other"


def analyze_companions(source_dir: Path) -> dict[str, Any]:
    json_paths = sorted(source_dir.glob("*.case.json"))
    matched = 0
    unreadable_json = 0
    unreadable_xlsx = 0
    unreadable_pdf = 0
    top_level_shapes: Counter[tuple[str, ...]] = Counter()
    section_shapes: Counter[tuple[str, ...]] = Counter()
    control_sets: list[set[str]] = []
    field_counts: list[int] = []
    unique_control_counts: list[int] = []
    field_count_frequencies: Counter[int] = Counter()
    control_set_frequencies: Counter[tuple[str, ...]] = Counter()
    section_field_counts: dict[str, list[int]] = {}
    duplicate_control_counts: list[int] = []
    type_by_control: dict[str, set[str]] = {}
    xlsx_sheets: list[int] = []
    xlsx_nonempty: list[int] = []
    xlsx_formulas: list[int] = []
    pdf_pages: list[int] = []
    pdf_pages_with_text: list[int] = []
    pdf_chars: list[int] = []
    xlsx_covered_by_json: list[float] = []
    json_covered_by_xlsx: list[float] = []
    pdf_covered_by_json: list[float] = []
    json_covered_by_pdf: list[float] = []
    pdf_covered_by_xlsx: list[float] = []
    xlsx_unique_token_counts: list[int] = []
    xlsx_cell_coverage: list[float] = []
    xlsx_cell_token_coverage: list[float] = []
    xlsx_unmatched_cell_counts: list[int] = []
    xlsx_unmatched_token_frequency: Counter[str] = Counter()
    xlsx_unmatched_cell_classes: Counter[str] = Counter()
    xlsx_unmatched_cell_frequency: Counter[str] = Counter()
    xlsx_unmatched_cell_class_by_value: dict[str, str] = {}
    xlsx_unmatched_cells_seen_in_pdf: list[float] = []
    pdf_unique_token_counts: list[int] = []
    pdf_unmatched_token_frequency: Counter[str] = Counter()

    for json_path in json_paths:
        stem = json_path.stem.removesuffix(".case")
        xlsx_path = source_dir / f"{stem}.xlsx"
        pdf_path = source_dir / f"{stem}.pdf"
        if not xlsx_path.exists() or not pdf_path.exists():
            continue
        matched += 1
        try:
            data = json.loads(json_path.read_text(encoding="utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            unreadable_json += 1
            continue
        top_level_shapes[tuple(sorted(data))] += 1
        sections = data.get("sections", {})
        section_shapes[tuple(sorted(sections)) if isinstance(sections, dict) else tuple()] += 1
        controls: list[str] = []
        if isinstance(sections, dict):
            for section_name, section in sections.items():
                if not isinstance(section, dict):
                    continue
                section_fields = section.get("fields", [])
                if not isinstance(section_fields, list):
                    section_fields = []
                section_field_counts.setdefault(section_name, []).append(len(section_fields))
                for field in section_fields:
                    if not isinstance(field, dict) or not isinstance(field.get("name"), str):
                        continue
                    name = field["name"]
                    controls.append(name)
                    input_type = field.get("type")
                    if isinstance(input_type, str):
                        type_by_control.setdefault(name, set()).add(input_type)
        control_sets.append(set(controls))
        field_counts.append(len(controls))
        unique_control_counts.append(len(set(controls)))
        field_count_frequencies[len(controls)] += 1
        control_set_frequencies[tuple(sorted(set(controls)))] += 1
        duplicate_control_counts.append(len(controls) - len(set(controls)))
        json_strings = _json_strings(data)
        json_tokens = _tokens(json_strings)
        json_blob = _normalized_scalar("\n".join(json_strings))
        unmatched_xlsx_cells: set[str] = set()

        try:
            xlsx_values, xlsx_metrics = _xlsx_strings(xlsx_path)
            xlsx_tokens = _tokens(xlsx_values)
            xlsx_sheets.append(xlsx_metrics["sheets"])
            xlsx_nonempty.append(xlsx_metrics["nonempty_cells"])
            xlsx_formulas.append(xlsx_metrics["formula_cells"])
            xlsx_covered_by_json.append(_coverage(xlsx_tokens, json_tokens))
            json_covered_by_xlsx.append(_coverage(json_tokens, xlsx_tokens))
            xlsx_unique_token_counts.append(len(xlsx_tokens - json_tokens))
            xlsx_unmatched_token_frequency.update(xlsx_tokens - json_tokens)
            normalized_cells = {
                normalized
                for value in xlsx_values
                if len(normalized := _normalized_scalar(value)) >= 2
            }
            matched_cells = {cell for cell in normalized_cells if cell in json_blob}
            xlsx_cell_coverage.append(_coverage(normalized_cells, matched_cells))
            unmatched_xlsx_cells = normalized_cells - matched_cells
            xlsx_unmatched_cell_frequency.update(unmatched_xlsx_cells)
            xlsx_unmatched_cell_counts.append(len(unmatched_xlsx_cells))
            original_by_normalized = {
                _normalized_scalar(value): value for value in xlsx_values
            }
            xlsx_unmatched_cell_classes.update(
                _unmatched_cell_class(original_by_normalized[cell], stem)
                for cell in unmatched_xlsx_cells
            )
            for cell in unmatched_xlsx_cells:
                xlsx_unmatched_cell_class_by_value[cell] = _unmatched_cell_class(
                    original_by_normalized[cell], stem
                )
            semantic_cells = [tokens for value in xlsx_values if (tokens := _tokens([value]))]
            semantically_covered_cells = [
                tokens for tokens in semantic_cells if tokens <= json_tokens
            ]
            xlsx_cell_token_coverage.append(
                len(semantically_covered_cells) / len(semantic_cells)
                if semantic_cells
                else 1.0
            )
        except Exception:  # aggregate audit records only the failure count
            unreadable_xlsx += 1
            xlsx_tokens = set()

        try:
            pdf_values, pdf_metrics = _pdf_strings(pdf_path)
            pdf_tokens = _tokens(pdf_values)
            pdf_blob = _normalized_scalar("\n".join(pdf_values))
            pdf_pages.append(pdf_metrics["pages"])
            pdf_pages_with_text.append(pdf_metrics["pages_with_text"])
            pdf_chars.append(pdf_metrics["extracted_characters"])
            pdf_covered_by_json.append(_coverage(pdf_tokens, json_tokens))
            json_covered_by_pdf.append(_coverage(json_tokens, pdf_tokens))
            pdf_covered_by_xlsx.append(_coverage(pdf_tokens, xlsx_tokens))
            pdf_unique_token_counts.append(len(pdf_tokens - json_tokens))
            pdf_unmatched_token_frequency.update(pdf_tokens - json_tokens)
            if unmatched_xlsx_cells:
                xlsx_unmatched_cells_seen_in_pdf.append(
                    sum(cell in pdf_blob for cell in unmatched_xlsx_cells)
                    / len(unmatched_xlsx_cells)
                )
        except Exception:  # aggregate audit records only the failure count
            unreadable_pdf += 1

    common_controls = set.intersection(*control_sets) if control_sets else set()
    union_controls = set.union(*control_sets) if control_sets else set()
    return {
        "analysis_kind": "aggregate-cross-format-structure-only",
        "contains_case_identifiers": False,
        "contains_source_values": False,
        "case_json_files": len(json_paths),
        "matched_json_xlsx_pdf": matched,
        "unreadable": {
            "json": unreadable_json,
            "xlsx": unreadable_xlsx,
            "pdf": unreadable_pdf,
        },
        "json": {
            "top_level_shape_variants": len(top_level_shapes),
            "section_shape_variants": len(section_shapes),
            "field_count_per_record": _distribution(field_counts),
            "unique_control_count_per_record": _distribution(unique_control_counts),
            "field_count_frequencies": {
                str(key): value for key, value in sorted(field_count_frequencies.items())
            },
            "control_set_variants": len(control_set_frequencies),
            "control_set_variant_record_counts": sorted(
                control_set_frequencies.values(), reverse=True
            ),
            "section_field_count_per_record": {
                section: _distribution(counts)
                for section, counts in sorted(section_field_counts.items())
            },
            "duplicate_controls_per_record": _distribution(duplicate_control_counts),
            "common_controls": len(common_controls),
            "union_controls": len(union_controls),
            "controls_with_input_type_drift": sum(
                len(input_types) > 1 for input_types in type_by_control.values()
            ),
        },
        "xlsx": {
            "sheets_per_workbook": _distribution(xlsx_sheets),
            "nonempty_cells_per_workbook": _distribution(xlsx_nonempty),
            "formula_cells_per_workbook": _distribution(xlsx_formulas),
        },
        "pdf": {
            "pages_per_document": _distribution(pdf_pages),
            "pages_with_extractable_text": _distribution(pdf_pages_with_text),
            "extracted_characters_per_document": _distribution(pdf_chars),
        },
        "token_coverage": {
            "xlsx_tokens_found_in_json": _distribution(xlsx_covered_by_json),
            "json_tokens_found_in_xlsx": _distribution(json_covered_by_xlsx),
            "pdf_tokens_found_in_json": _distribution(pdf_covered_by_json),
            "json_tokens_found_in_pdf": _distribution(json_covered_by_pdf),
            "pdf_tokens_found_in_xlsx": _distribution(pdf_covered_by_xlsx),
            "xlsx_tokens_not_found_in_json": _distribution(xlsx_unique_token_counts),
            "xlsx_nonempty_cells_found_in_json": _distribution(xlsx_cell_coverage),
            "xlsx_cells_with_all_tokens_found_in_json": _distribution(
                xlsx_cell_token_coverage
            ),
            "xlsx_nonempty_cells_not_found_in_json": _distribution(
                xlsx_unmatched_cell_counts
            ),
            "xlsx_unmatched_cell_classes": dict(
                sorted(xlsx_unmatched_cell_classes.items())
            ),
            "xlsx_unmatched_cell_document_frequency": {
                "distinct": len(xlsx_unmatched_cell_frequency),
                "one_record": sum(
                    count == 1 for count in xlsx_unmatched_cell_frequency.values()
                ),
                "two_to_five_records": sum(
                    2 <= count <= 5 for count in xlsx_unmatched_cell_frequency.values()
                ),
                "more_than_five_records": sum(
                    count > 5 for count in xlsx_unmatched_cell_frequency.values()
                ),
                "all_records": sum(
                    count == matched for count in xlsx_unmatched_cell_frequency.values()
                ),
            },
            "xlsx_one_record_unmatched_cells_by_class": dict(
                sorted(
                    Counter(
                        xlsx_unmatched_cell_class_by_value[value]
                        for value, count in xlsx_unmatched_cell_frequency.items()
                        if count == 1
                    ).items()
                )
            ),
            "xlsx_unmatched_cells_also_found_in_pdf": _distribution(
                xlsx_unmatched_cells_seen_in_pdf
            ),
            "pdf_tokens_not_found_in_json": _distribution(pdf_unique_token_counts),
            "xlsx_unmatched_token_document_frequency": {
                "distinct": len(xlsx_unmatched_token_frequency),
                "one_record": sum(
                    count == 1 for count in xlsx_unmatched_token_frequency.values()
                ),
                "two_to_five_records": sum(
                    2 <= count <= 5 for count in xlsx_unmatched_token_frequency.values()
                ),
                "more_than_five_records": sum(
                    count > 5 for count in xlsx_unmatched_token_frequency.values()
                ),
            },
            "pdf_unmatched_token_document_frequency": {
                "distinct": len(pdf_unmatched_token_frequency),
                "one_record": sum(
                    count == 1 for count in pdf_unmatched_token_frequency.values()
                ),
                "two_to_five_records": sum(
                    2 <= count <= 5 for count in pdf_unmatched_token_frequency.values()
                ),
                "more_than_five_records": sum(
                    count > 5 for count in pdf_unmatched_token_frequency.values()
                ),
            },
        },
    }
