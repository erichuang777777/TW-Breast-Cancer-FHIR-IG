from __future__ import annotations
import json,re
from pathlib import Path
import fitz
from openpyxl import load_workbook
from .models import BatchRecord,CaseRecord,Method,TreatmentEvent
from .rules import derive_stage,ev,extract_pathology,mark_applicability,put

ALL_D=[f"D{i:03d}" for i in range(1,86)]

def _control_key(field):
    parts=str(field.get("name","")).split("$")
    return parts[-2] if len(parts)>1 and parts[-1].isdigit() else parts[-1]

def digits_date(value):
    nums=re.findall(r"\d+",str(value or ""));
    if len(nums)>=3: return f"{int(nums[0]):04d}{int(nums[1]):02d}{int(nums[2]):02d}"
    return None

def selected(fields,suffix):
    hits=[f for f in fields if _control_key(f)==suffix]
    for f in hits:
        if f.get("checked") is True: return f.get("label") or f.get("selected_text") or f.get("value"),f.get("name")
    for f in hits:
        v=f.get("selected_text") or f.get("value")
        if v not in (None,"","請選擇"): return v,f.get("name")
    return None,None

def selected_all(fields,suffix):
    answers=[]
    for field in fields:
        if _control_key(field)!=suffix or field.get("checked") is not True:
            continue
        value=field.get("label") or field.get("selected_text") or field.get("value")
        if value not in (None,"","請選擇"):
            answers.append((str(value),field.get("name")))
    return answers

def table_value(data,label):
    section=data.get("sections",{}).get("basic",{})
    wanted=str(label).strip().rstrip("：:")
    for table_index,table in enumerate(section.get("tables",[])):
        for row_index,row in enumerate(table.get("rows",[])):
            cells=[str(value).strip() for value in row if str(value).strip()]
            if len(cells)>=2 and cells[0].rstrip("：:")==wanted:
                return cells[1],f"sections.basic.tables[{table_index}].rows[{row_index}]"
    return None,None

def import_case_json(path:Path,case:CaseRecord):
    data=json.loads(path.read_text(encoding="utf-8")); fields=data["sections"]["basic"].get("fields",[]); filename=path.name
    def add(tag,suffix,mapping=None,display=True,review=False):
        v,p=selected(fields,suffix)
        if v is None:return
        q=mapping.get(str(v),str(v)) if mapping else str(v)
        put(case,tag,q,Method.STRUCTURED,ev(filename,"json",path=p),"JSON_FIELD_MAP",str(v) if display else None,review=review)
    case.source_files.append(filename)
    for tag,label in [("P01","姓名"),("BIRTHDAY","生日")]:
        v,p=table_value(data,label)
        if v:
            value=digits_date(v) if tag=="BIRTHDAY" else v
            if value: put(case,tag,value,Method.STRUCTURED,ev(filename,"json_table",path=p),"JSON_TABLE_DEMOGRAPHIC")
    sex,sex_path=table_value(data,"性別")
    sex_map={"M":"0","男":"0","男性":"0","F":"1","女":"1","女性":"1","其他":"2","未知":"3"}
    if str(sex).strip() in sex_map:
        put(case,"P02",sex_map[str(sex).strip()],Method.STRUCTURED,ev(filename,"json_table",path=sex_path),"JSON_TABLE_SEX")
    add("DIAG_TYPE","ddlReason",{"初診斷或初次治療":"2"})
    case.diagnosis_type=case.candidates.get("DIAG_TYPE").value if case.candidates.get("DIAG_TYPE") else "2"
    add("LATERALITY","rblLocation",{"左側":"L","右側":"R"}); case.laterality=case.candidates.get("LATERALITY").value if case.candidates.get("LATERALITY") else None
    add("P05","rblMenopause",{"是":"1","否":"2"},review=True)
    add("D002","rblHospital",{"本院":"1","外院":"2"})
    histology_map={
        "Ductal carcinoma in situ":"1",
        "Invasive ductal carcinoma":"2",
        "Infiltation ductal carcinoma":"2",
        "Invasive lobular carcinoma":"3",
        "Infiltation lobular carcinoma":"3",
        "Mucinous carcinoma":"4",
        "Other":"8",
        "Phyllodes tumor, malignant":"8",
    }
    histology_answers=selected_all(fields,"cblHisType")
    histology_codes={histology_map[value] for value,_ in histology_answers if value in histology_map}
    if case.diagnosis_type=="2" and len(histology_codes)==1:
        histology_code=next(iter(histology_codes))
        source_path=next(path for value,path in histology_answers if histology_map.get(value)==histology_code)
        put(case,"D003",histology_code,Method.STRUCTURED,ev(filename,"json",path=source_path),"JSON_HISTOLOGY_TYPE",review=True)
    elif case.diagnosis_type=="2" and len(histology_codes)>1:
        case.issues.append("D003 has multiple histology categories and requires review")
    mapping=[("D004","rb2HGrade1",{"Ⅰ":"1","Ⅱ":"2","Ⅲ":"3","Unknown":"X"}),("D018","rb2Her1",{"Unknown":"X","0+":"0","1+":"1","2+":"2","3+":"3"}),("D020","rbl2Ki67",{"未檢測":"0","已檢測：":"1"}),("D021","txb2Ki67",None),("D031","rblHGrade",{"Ⅰ":"1","Ⅱ":"2","Ⅲ":"3","Unknown":"X"}),("D052","rblHer",{"Unknown":"X","0+":"0","1+":"1","2+":"2","3+":"3"}),("D054","rblKi67",{"未檢測":"0","已檢測：":"1"}),("D055","txbKi67",None),("D047","txbSize1",None)]
    for x in mapping:add(*x)
    fish_mappings = [
        ("D019", "rb2HerFISH1", "D018"),
        ("D053", "rblHerFISH", "D052"),
    ]
    for tag, suf, ihc_tag in fish_mappings:
        v, p = selected(fields, suf)
        # Source options: 0=not tested, 1=positive, 2=negative. QBC only
        # accepts an actual result here: 1=positive, 0=negative.
        positive = {"1", "Positive", "positive", "陽性"}
        negative = {"2", "Negative", "negative", "陰性"}
        has_equivocal_ihc = (
            case.candidates.get(ihc_tag)
            and case.candidates[ihc_tag].value == "2"
        )
        if str(v) in positive | negative and has_equivocal_ihc:
            put(
                case,
                tag,
                "1" if str(v) in positive else "0",
                Method.STRUCTURED,
                ev(filename, "json", path=p),
                "JSON_HER2_FISH",
            )
        elif str(v) in positive | negative:
            case.issues.append(
                f"{tag} source result ignored because {ihc_tag} is not equivocal"
            )
    for tag,suf in [("D015","txb2EReceptor1"),("D017","txb2PReceptor1"),("D049","txbEReceptor"),("D051","txbPReceptor")]:
        v,_=selected(fields,suf)
        put(case,tag,str(v),Method.STRUCTURED,ev(filename,"json",path=suf),"JSON_FIELD_MAP") if v not in (None,"0") else None
    for tag,suf in [("D014","txb2EReceptor1"),("D016","txb2PReceptor1"),("D048","txbEReceptor"),("D050","txbPReceptor")]:
        v,_=selected(fields,suf)
        if v is not None: put(case,tag,"0" if str(v)=="0" else "1",Method.RULE_DERIVED,ev(filename,"json",path=suf),"RECEPTOR_VALUE_TO_STATUS")
    for tag,suf,prefix in [("D005","ddlClinicTGeneral","T"),("D006","ddlClinicNGeneral","N"),("D007","ddlClinicMGeneral","M"),("D032","ddlPathTGeneral","T"),("D033","ddlPathNGeneral","N"),("D034","ddlPathMGeneral","M")]:
        v,p=selected(fields,suf)
        if v is not None: put(case,tag,prefix+str(v),Method.STRUCTURED,ev(filename,"json",path=p),"JSON_TNM")
    derive_stage(case,"D005","D006","D007","D008"); derive_stage(case,"D032","D033","D034","D035")
    margin,margin_path=selected(fields,"rblMargin")
    if margin=="Negative":
        put(case,"D028","0",Method.STRUCTURED,ev(filename,"json",path=margin_path),"JSON_SURGICAL_MARGIN",review=True)
    elif margin=="Positive":
        case.issues.append("D028 positive margin requires DCIS versus invasive review")
    metastasis_map={
        "LN-distant":"1","Bone":"3","Liver":"4","Lung":"5","Brain":"6",
        "Pleural seeding":"8","Peritoneal seeding":"9","Adrenal":"10",
    }
    metastasis_answers=selected_all(fields,"cblMetastasis")
    metastasis_codes={metastasis_map.get(value,"12") for value,_ in metastasis_answers}
    other_sites=sorted(value for value,_ in metastasis_answers if value not in metastasis_map)
    for stage_tag,site_tag,other_tag in [("D008","D009","D010"),("D035","D036","D037")]:
        if case.candidates.get(stage_tag) and case.candidates[stage_tag].value=="StageⅣ" and metastasis_codes:
            source_path=metastasis_answers[0][1]
            put(case,site_tag,",".join(sorted(metastasis_codes,key=int)),Method.STRUCTURED,ev(filename,"json",path=source_path),"JSON_METASTASIS_SITES",review=True)
            if "12" in metastasis_codes and other_sites:
                put(case,other_tag,", ".join(other_sites),Method.STRUCTURED,ev(filename,"json",path=source_path),"JSON_METASTASIS_OTHER",review=True)
    mark_applicability(case,ALL_D)
    reports=str(data["sections"].get("reference_reports",{}).get("text","")); extract_pathology(case,reports,filename)
    plan=str(data["sections"].get("treatment_plan",{}).get("text",""))
    for date,kind in re.findall(r"填表日期[：:]\s*(\d{4}/\d{1,2}/\d{1,2}).{0,120}?\[(抗癌治療|手術)\]",plan,re.S):
        case.issues.append(f"planned treatment candidate {kind} {date}; verify against execution record")

def read_eligible(path:Path):
    raw=path.read_bytes(); text=raw.decode("cp950",errors="replace")
    import csv,io
    return list(csv.DictReader(io.StringIO(text),delimiter="\t"))

def apply_eligible(row,case:CaseRecord,filename):
    evidence=ev(filename,"eligible_list")
    maps={"P09":row.get("收案日"),"D001":row.get("診斷年月"),"P08":row.get("Class分類")}
    for tag,v in maps.items():
        val=digits_date(v) if tag in {"P09","D001"} else v
        if val: put(case,tag,val,Method.STRUCTURED,evidence,"ELIGIBLE_COLUMN_MAP")
    for key,tcode in [("手術","1"),("化療","2"),("放療","3")]:
        text=row.get(key) or ""
        if not text.strip(): continue
        dates=[digits_date(x) for x in re.findall(r"20\d{2}/\d{1,2}/\d{1,2}",text)]
        doctors=re.findall(r"<主治醫師>:\s*([^\n]+)",text)
        case.treatments.append(TreatmentEvent(sequence=len(case.treatments)+1,treatment_type=tcode,
            actual_start=dates[0] if dates else None,actual_end=dates[-1] if len(dates)>1 else None,
            evidence=[ev(filename,"eligible_list",text=text[:1000],path=key)]))

def extract_document_text(path:Path):
    if path.suffix.lower()==".pdf": return "\n".join(p.get_text() for p in fitz.open(path))
    if path.suffix.lower()==".docx":
        from docx import Document
        d=Document(path); return "\n".join([p.text for p in d.paragraphs]+[c.text for t in d.tables for r in t.rows for c in r.cells])
    if path.suffix.lower()==".xlsx":
        wb=load_workbook(path,data_only=True); return "\n".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value not in (None,""))
    return ""

def import_batch(source:Path,batch_id:str,case_ids:list[str]|None=None):
    batch=BatchRecord(batch_id=batch_id,source_directory=str(source)); eligible=None
    for p in source.glob("EligibleList*.xls"):
        eligible=read_eligible(p); eligible_name=p.name; break
    jsons=list(source.rglob("*.case.json"))
    for jp in jsons:
        cid=jp.name.split(".")[0]
        if case_ids and cid not in case_ids: continue
        case=CaseRecord(case_id=cid); import_case_json(jp,case)
        if eligible:
            for row in eligible:
                if str(row.get("病歷號","")).strip()==cid: apply_eligible(row,case,eligible_name); break
        for ext in (".pdf",".docx",".xlsx"):
            for p in source.rglob(cid+ext):
                txt=extract_document_text(p); case.source_files.append(p.name)
                if "patholog" in txt.lower() or "病理" in txt: extract_pathology(case,txt,p.name)
        batch.cases[cid]=case
    return batch
